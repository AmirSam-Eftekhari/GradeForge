from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.models.domain import GradingResult

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def write_class_excel_report(results: list[GradingResult], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Class Report"

    question_numbers: list[str] = []
    for r in results:
        for fb in r.per_question:
            if fb.question_number not in question_numbers:
                question_numbers.append(fb.question_number)

    headers = ["Student Name", "Student ID", "Class"] + [f"Q{q}" for q in question_numbers] + [
        "Total Score", "Max Score", "Percentage",
    ]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for result in results:
        scores_by_q = {fb.question_number: fb.score for fb in result.per_question}
        row = [
            result.student.name or "(unidentified)",
            result.student.student_id or "",
            result.student.class_name or "",
        ]
        row += [scores_by_q.get(q, "") for q in question_numbers]
        row += [result.total_score, result.max_total_score, f"{result.percentage:.1f}%"]
        ws.append(row)

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

    ws.freeze_panes = "A2"
    wb.save(str(output_path))
