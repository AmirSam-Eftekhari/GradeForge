"""
Reads raw text out of any supported input format.

Digital-native formats (.docx, .txt) are read directly — running OCR on
text that is already machine-readable would be a needless accuracy hit.
Everything else (PDF, scanned PDF, PNG/JPEG/TIFF) goes through the OCR
engine. The distinction between "digital PDF" and "scanned PDF" is made
automatically: we first try to extract an embedded text layer, and only
fall back to OCR if that layer is empty/near-empty.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from src.ocr.base import OCREngine


@dataclass
class ExtractedDocument:
    text: str
    confidence: float  # 1.0 for digital-native text; OCR mean-confidence otherwise
    used_ocr: bool


def _read_docx(file_path: Path) -> str:
    import docx  # python-docx

    doc = docx.Document(str(file_path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _read_txt(file_path: Path) -> str:
    return file_path.read_text(encoding="utf-8", errors="replace")


def _try_pdf_text_layer(file_path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    try:
        reader = PdfReader(str(file_path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:
        return ""


def read_document(file_path: Path, ocr_engine: OCREngine, languages: list[str]) -> ExtractedDocument:
    suffix = file_path.suffix.lower()

    if suffix == ".docx":
        return ExtractedDocument(text=_read_docx(file_path), confidence=1.0, used_ocr=False)
    if suffix == ".txt":
        return ExtractedDocument(text=_read_txt(file_path), confidence=1.0, used_ocr=False)

    if suffix == ".pdf":
        text_layer = _try_pdf_text_layer(file_path)
        if len(text_layer.strip()) > 40:  # heuristic: "real" digital text, not noise
            return ExtractedDocument(text=text_layer, confidence=1.0, used_ocr=False)
        ocr_doc = ocr_engine.extract(file_path, languages)
        return ExtractedDocument(text=ocr_doc.full_text, confidence=ocr_doc.mean_confidence, used_ocr=True)

    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}:
        ocr_doc = ocr_engine.extract(file_path, languages)
        return ExtractedDocument(text=ocr_doc.full_text, confidence=ocr_doc.mean_confidence, used_ocr=True)

    raise ValueError(f"Unsupported file type: {suffix}")
